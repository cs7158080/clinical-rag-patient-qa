"""
docx_factory.py — synthetic .docx builders for parser tests (PLAN.md section 7, T0 part 5).

Content is invented; structure mirrors the client's corpus. Files are written
to the caller's tmp_path — never committed, never read from the corpus.
"""

from dataclasses import dataclass

from docx import Document


@dataclass
class Table:
    """A Word table. Each row is a list of cell texts; a row with ONE text is
    merged across all columns (title / patient-name / note row, as in the corpus)."""
    rows: list[list[str]]


def write_docx(path, items: list):
    """Write *items* in order and return *path*: a str is one Word paragraph,
    a Table is a table.

    A '\\n' inside a paragraph or a cell becomes a Shift+Enter line break,
    as therapists type it in Word.
    """
    doc = Document()
    for item in items:
        if isinstance(item, Table):
            _add_table(doc, item)
            continue
        para = doc.add_paragraph()
        for i, line in enumerate(item.split("\n")):
            run = para.add_run(line)
            if i < len(item.split("\n")) - 1:
                run.add_break()
    doc.save(path)
    return path


def _add_table(doc, table: Table):
    width = max(len(row) for row in table.rows)
    docx_table = doc.add_table(rows=len(table.rows), cols=width)
    for docx_row, row in zip(docx_table.rows, table.rows):
        cells = docx_row.cells
        if len(row) == 1 and width > 1:
            cells[0].merge(cells[-1]).text = row[0]
        else:
            for cell, text in zip(cells, row):
                cell.text = text
